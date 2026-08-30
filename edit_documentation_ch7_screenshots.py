# -*- coding: utf-8 -*-
r"""
Fill the three "Insert Screenshot" placeholders in GarageAI_Documentation.docx
Chapter 7 (User Guide) with real captures of the running app, taken
2026-08-31 against the full local stack (frontend 5173 + gateway 5000 +
FastAPI inference service 8001) via Playwright.

Each placeholder is a 1x1 bordered table holding a "img  Insert Screenshot: ..."
prompt; this replaces the cell content with the image (keeping the border as
a figure frame) and adds a centered caption paragraph beneath the table.

  7.1 Recording or Uploading Audio     -> guide_7_1_home.png
  7.2 Viewing the Classification Result -> guide_7_2_result.png
  7.3 Changing the Display Theme        -> guide_7_3_theme.png
"""
import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches
from docx.table import Table
from docx.text.paragraph import Paragraph

DOC = r"C:\Users\hp\Downloads\GarageAI_Documentation.docx"
SHOT_DIR = r"C:\Users\hp\AppData\Local\Temp\claude\c--Users-hp-Desktop-car-test\f64fe89b-64da-445d-b0e5-e0e0c3213a35\scratchpad"

SHOTS = [
    ("From the Home page, the user can either tap the microphone button",
     rf"{SHOT_DIR}\guide_7_1_home.png", 5.9,
     "Figure 7.1. The Home page \u2014 record a clip directly in the browser or upload an "
     "existing audio file, with an optional extra-model comparison selector."),
    ("After submission, the app displays the predicted fault class",
     rf"{SHOT_DIR}\guide_7_2_result.png", 5.9,
     "Figure 7.2. A returned classification \u2014 predicted class and overall confidence, "
     "the per-class confidence breakdown, and each individual model's vote."),
    ("The Settings page (and the quick toggle in the navigation bar)",
     rf"{SHOT_DIR}\guide_7_3_theme.png", 5.4,
     "Figure 7.3. The Appearance section of the Settings page, with the Light / Dark / "
     "System theme selector."),
]


def main():
    d = docx.Document(DOC)
    P = d.paragraphs

    cap_style = None
    for p in P:
        s = p.text.strip()
        if s.startswith("Figure 7.") or s.startswith("Figure 4.") or s.startswith("Figure 3."):
            cap_style = p.style
            break

    for anchor_prefix, img_path, width_in, caption in SHOTS:
        body = next(p for p in d.paragraphs if p.text.startswith(anchor_prefix))
        tbl_el = body._p.getnext()
        assert tbl_el is not None and tbl_el.tag == qn("w:tbl"), \
            f"expected a table after {anchor_prefix!r}, got {tbl_el}"
        tbl = Table(tbl_el, body._parent)
        cell = tbl.rows[0].cells[0]

        assert "Insert Screenshot" in cell.text, f"unexpected cell content: {cell.text!r}"
        cell.text = ""  # collapse to a single empty paragraph
        cpar = cell.paragraphs[0]
        cpar.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cpar.add_run().add_picture(img_path, width=Inches(width_in))

        # caption paragraph immediately after the table
        new_p = d.add_paragraph()  # appended at end of body; move it after the table
        tbl_el.addnext(new_p._p)
        if cap_style is not None:
            new_p.style = cap_style
        new_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        new_p.add_run(caption)

    d.save(DOC)
    print("saved", DOC)


if __name__ == "__main__":
    main()
