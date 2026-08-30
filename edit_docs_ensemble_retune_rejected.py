# -*- coding: utf-8 -*-
r"""
Document the 2026-08-31 negative result: a nested-CV retune of the fused
ensemble weights around the fine-tuned EfficientAT model
({Traditional 0.14, CNN 0.18, PANNs 0.04, EfficientAT 0.64}) raised CLEAN
test accuracy 89.47% -> 90.23% (+0.76 pt) but regressed the fused result
to 76.2% under the real WebM/Opus production pipeline (vs 84.0% for the
deployed {0.30/0.60/0.10} config on identical audio; Brake 84.1% -> 45.5%
at 32 kbps). Cause: EfficientAT, which takes the dominant 0.64 weight, was
never compression-augmented. The retune was rejected; the original
3-model config stays deployed.

Evidence:
  car_test/nested_cv_ensemble_5model_alwayson.py
  car_test/mic_pipeline_test/test_full_ensemble_compression_4model.py
  car_test/mic_pipeline_test/full_ensemble_compression_4model_results.csv

Touches both deliverables (research paper + documentation, whose
Appendix A reproduces the paper). Adds a new subsection 3.4.6 and updates
the Discussion (7.2), Limitations, Future Work, and Conclusion wording
that previously framed the EfficientAT question as open.
"""
import docx

PAPER = r"C:\Users\hp\Downloads\GarageAI_Research_Paper.docx"
DOCUMENTATION = r"C:\Users\hp\Downloads\GarageAI_Documentation.docx"

SUBSEC_HEADING = "3.4.6 A Post-Hoc Retune, Tested and Rejected"

SUBSEC_BODY = (
    "The open question of whether folding a stronger backbone into the fused weights "
    "would improve the deployed result was tested directly. A nested five-fold "
    "cross-validation weight search, restricted to the five models already computed on "
    "every prediction (Traditional ML, CNN, YAMNet, PANNs/CNN14, and the fine-tuned "
    "EfficientAT), produced a more stable out-of-fold macro-F1 than the deployed "
    "three-model configuration (0.907 versus 0.860); averaging the five folds' winning "
    "weight vectors gave Traditional ML 0.14 / CNN 0.18 / PANNs/CNN14 0.04 / EfficientAT "
    "0.64. Evaluated once on the untouched test set, this raised clean accuracy from "
    "89.47% to 90.23% (+0.76 points). It was nonetheless not adopted. Re-evaluated "
    "end-to-end through the real WebM/Opus production pipeline (Section 3.3.2) on all "
    "133 test recordings, the retuned ensemble regressed sharply under compression -- "
    "76.2% mean accuracy across 16/32/64 kbps, against 84.0% for the deployed "
    "configuration on the identical audio -- with Brake collapsing from 84.1% clean to "
    "45.5% at 32 kbps. Per-model inspection traced the regression to EfficientAT, which "
    "holds the dominant 0.64 weight in the retuned configuration but, unlike the "
    "Traditional ML and CNN models, was never compression-augmented: it is confidently "
    "wrong on compressed audio and overrides the other three models even when they "
    "agree on the correct class -- the same recording-chain-artifact failure mode "
    "diagnosed in Section 3.3.2. For a system whose live input is always compressed, "
    "the small clean-accuracy gain is a poor trade, and the original 0.30 / 0.60 / 0.10 "
    "Traditional ML / CNN / YAMNet ensemble was retained."
)

# --- 7.2 Discussion: the paragraph that previously left the question open ---
DISC_OLD_TAIL = (
    "and raises a natural question \u2014 taken up as future work in Section 9 \u2014 of "
    "whether re-tuning the ensemble to include AST or EfficientAT would push this "
    "further, or whether their errors are already well-covered by the existing three "
    "models."
)
DISC_NEW_TAIL = (
    "and raised a natural question of whether re-tuning the ensemble to include a "
    "stronger backbone such as EfficientAT would push this further. Section 3.4.6 "
    "reports the test: a nested cross-validation retune giving EfficientAT the dominant "
    "weight improved clean test accuracy by 0.76 points but regressed the fused result "
    "by roughly eight points under the production compression pipeline -- EfficientAT "
    "never having been compression-augmented -- and was not adopted."
)

# --- Limitations bullet (paper P177 / documentation Appendix-A P331) ---
LIM_OLD = (
    "The production ensemble does not yet include AST or EfficientAT, the two strongest "
    "individual models identified in this extended comparison; this is a deliberate "
    "decision pending further evaluation, not an oversight (Section 9)."
)
LIM_NEW = (
    "The production ensemble does not include AST or EfficientAT, the two strongest "
    "individual models identified in this extended comparison. Folding EfficientAT in "
    "via a cross-validated weight retune was tested (Section 3.4.6): it improved clean "
    "test accuracy by 0.76 points (89.47% to 90.23%) but regressed the ensemble to "
    "76.2% under the real WebM/Opus compression pipeline, against 84.0% for the "
    "deployed configuration, and was rejected because that model was never "
    "compression-augmented."
)

# --- Future Work bullet (paper P191 / documentation Appendix-A P345) ---
FW_OLD = (
    "Make an explicit decision on whether to re-tune the production ensemble to include "
    "AST or EfficientAT, given their clear individual improvement over the models "
    "currently in the ensemble."
)
FW_NEW = (
    "Compression-augment EfficientAT (and AST) during fine-tuning so that a stronger "
    "backbone can be folded into the production ensemble without the compression "
    "regression documented in Section 3.4.6."
)

# --- Conclusion sentence (paper P196 / documentation Appendix-A P350) ---
CONC_ANCHOR = (
    "remained the strongest configuration overall at 89.47% accuracy, ahead of every "
    "individual model including AST and EfficientAT."
)
CONC_ADDED = (
    " A later attempt to fold the fine-tuned EfficientAT into the fused weights via "
    "cross-validated retuning was tested and rejected: it improved clean accuracy by "
    "0.76 points but lost roughly eight points under the application's real WebM/Opus "
    "compression pipeline, that model never having been compression-augmented "
    "(Section 3.4.6)."
)


def set_text(paragraph, new_text):
    if not paragraph.runs:
        paragraph.add_run(new_text)
        return
    paragraph.runs[0].text = new_text
    for r in paragraph.runs[1:]:
        r.text = ""


def replace_in_paragraph(paragraph, old, new):
    assert old in paragraph.text, f"anchor not found: {old[:60]!r}\nin: {paragraph.text[:200]!r}"
    set_text(paragraph, paragraph.text.replace(old, new))


def insert_after(anchor, text, style_obj=None):
    new_p = anchor.insert_paragraph_before("")
    anchor._p.addnext(new_p._p)
    if style_obj is not None:
        new_p.style = style_obj
    new_p.add_run(text)
    return new_p


def find_para(P, predicate, label):
    for p in P:
        if predicate(p.text):
            return p
    raise AssertionError(f"could not locate paragraph: {label}")


def common_edits(d, is_appendix_a):
    """Edits that apply to the paper AND to the documentation's Appendix-A copy."""
    P = d.paragraphs

    sway = find_para(
        P,
        lambda t: t.startswith("The remaining, largest source of ensemble error is Sway"),
        "3.4.5 sway paragraph",
    )
    heading_style = find_para(
        P, lambda t: t.startswith("3.4.5 Ensemble Selection"), "3.4.5 heading"
    ).style
    # insert body first, then the heading above it (addnext stacking)
    insert_after(sway, SUBSEC_BODY)
    insert_after(sway, SUBSEC_HEADING, style_obj=heading_style)

    disc = find_para(P, lambda t: DISC_OLD_TAIL in t, "7.2 discussion tail")
    replace_in_paragraph(disc, DISC_OLD_TAIL, DISC_NEW_TAIL)

    lim = find_para(P, lambda t: t.strip() == LIM_OLD, "limitations ensemble bullet")
    set_text(lim, LIM_NEW)

    fw = find_para(P, lambda t: t.strip() == FW_OLD, "future-work ensemble bullet")
    set_text(fw, FW_NEW)

    conc = find_para(P, lambda t: CONC_ANCHOR in t and "This research evaluated" in t,
                     "conclusion paragraph")
    replace_in_paragraph(conc, CONC_ANCHOR, CONC_ANCHOR + CONC_ADDED)


def edit_paper():
    d = docx.Document(PAPER)
    common_edits(d, is_appendix_a=False)
    d.save(PAPER)
    print("saved", PAPER)


def edit_documentation():
    d = docx.Document(DOCUMENTATION)
    P = d.paragraphs

    # Appendix A copy: same structural edits as the paper
    common_edits(d, is_appendix_a=True)

    # Chapter 4.4 "Deployed Model" -- summary sentence
    p87 = find_para(
        P,
        lambda t: t.startswith("The system's validation-set weight search favors the CNN"),
        "Ch 4.4 deployed-model paragraph",
    )
    tail = " Full derivation, statistical caveats, and discussion of this result are given in Appendix A, Sections 3.4.4 and 7.2."
    assert tail in p87.text
    set_text(p87, p87.text.replace(
        tail,
        tail + " A later cross-validated attempt to re-tune these weights around the "
        "fine-tuned EfficientAT model raised clean test accuracy to 90.23% but regressed "
        "the ensemble to 76.2% under the real WebM/Opus compression pipeline "
        "(EfficientAT was never compression-augmented) and was rejected; the "
        "0.30 / 0.60 / 0.10 configuration remains deployed (Appendix A, Section 3.4.6).",
    ))

    # Chapter 8.2 Limitations (main-body copy, worded slightly differently)
    p145 = find_para(
        P,
        lambda t: t.startswith("The production ensemble does not yet include AST or EfficientAT")
        and "Computational benchmarking" in t,
        "Ch 8.2 limitations ensemble bullet",
    )
    set_text(p145,
        "The production ensemble does not include AST or EfficientAT, the two strongest "
        "individual models identified in the extended comparison (Appendix A, Section 6). "
        "Folding EfficientAT in via a cross-validated weight retune was tested and "
        "rejected (Appendix A, Section 3.4.6): it raised clean accuracy by 0.76 points "
        "but regressed the ensemble to 76.2% under the real WebM/Opus compression "
        "pipeline. Computational benchmarking (latency, memory, model size) has not yet "
        "been performed for any of the ten configurations.")

    # Chapter 8.3 Future Work (main-body copy)
    p148 = find_para(
        P,
        lambda t: t.startswith("Re-verify and document the Weighted Ensemble's exact test-set accuracy"),
        "Ch 8.3 future-work regression-check bullet",
    )
    set_text(p148,
        "Add automated regression checks that re-verify the deployed ensemble's test-set "
        "accuracy -- and its accuracy through the real WebM/Opus compression pipeline -- "
        "whenever any underlying model weight changes, so that a clean-accuracy-only "
        "gain masking a compression regression (Appendix A, Section 3.4.6) is caught "
        "automatically.")

    # Appendix-A future-work note about folding models into the fused weights
    p344 = find_para(
        P,
        lambda t: "the remaining open question is whether any of them should move from opt-in to always-on" in t,
        "Appendix-A future-work fold-in note",
    )
    assert p344.text.rstrip().endswith("(Section 6.3, Table 7).")
    set_text(p344, p344.text.rstrip()[:-1] +
             "; a first attempt to fold EfficientAT into the fused ensemble weights was "
             "tested and rejected for a compression regression (Section 3.4.6).")

    d.save(DOCUMENTATION)
    print("saved", DOCUMENTATION)


if __name__ == "__main__":
    edit_paper()
    edit_documentation()
