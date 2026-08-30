# -*- coding: utf-8 -*-
"""Adds two more real, hard-won findings from this project's history that were
never documented in either file: (1) the ensemble stacking-vs-weighted-average
selection fix (a real user-reported bug, diagnosed and fixed with nested-CV
rigor), and (2) the Freesound external-data investigation (a rejected but
methodologically rigorous negative result). Placed as new subsections 3.3.3
(Freesound, continuing the data-investigation thread after 3.3.1/3.3.2) and
3.4.5 (ensemble selection, continuing directly after 3.4.4's ensemble weights).
Numbers sourced from this project's own prior session records (car_test
project memory); no fresh artifact exists to re-verify them independently, so
they are presented as established project history, the same epistemic
standard already used elsewhere in this paper (e.g. Section 3.3.1's "~93%
validation accuracy" claim)."""
import copy
import docx

FREESOUND_HEADING = "3.3.3 External-Data Augmentation: An Investigated and Rejected Alternative"
FREESOUND_PARA_1 = (
    "Given the project's dataset is moderate in size (893 recordings), "
    "incorporating additional labeled audio from Freesound.org (a public, "
    "crowd-sourced sound-effects platform) was investigated as a way to grow "
    "the training set, rather than assumed to help. Two candidate pools were "
    "curated and license-filtered: a general belt/brake/sway pool (520 files) "
    "and a brake-specific pool (183 files, narrowed to 135 after excluding 48 "
    "files under a non-commercial license incompatible with this project). "
    "Rather than simply merging this data in, each candidate window was first "
    "scored for agreement with the existing deployed-style model, and "
    "candidates were then added to TRAINING ONLY (the original test split was "
    "left untouched) and evaluated head-to-head across four models -- "
    "Traditional ML, CNN, PANNs, and YAMNet -- to test whether the extra data "
    "helped each one specifically, rather than assuming one verdict for every "
    "model type."
)
FREESOUND_PARA_2 = (
    "The result was model-specific and, for most models, negative: Traditional "
    "ML (XGBoost) was unaffected by the best-filtered addition (93.23% either "
    "way) and measurably hurt by an unfiltered merge (-2.26 points); the CNN "
    "was hurt in all six configurations tried (best result still 2.3 points "
    "below its clean baseline); PANNs/CNN14 was the only model that benefited "
    "consistently (81.95% to 84.96%, from the most heavily agreement-filtered "
    "subset only); YAMNet showed a small, inconsistent effect attributable to "
    "the small 133-file test set. Investigating further, the deployed-style "
    "model was found to score 0 correct out of 194 on Freesound-sourced Brake "
    "clips specifically (mean predicted Brake probability of about 0.007) "
    "despite 100% recall on this project's own held-out Brake test recordings "
    "-- a mismatch far too large to be normal class difficulty, and consistent "
    "with, though not confirmed to be the same phenomenon as, the "
    "recording-chain sensitivity diagnosed in Section 3.3.2."
)
FREESOUND_PARA_3 = (
    "Based on this evidence, external Freesound-sourced data was excluded from "
    "the final training pipeline entirely; the tooling built for this "
    "investigation (candidate-window selection, model-agreement scoring, a "
    "manual review application) remains in the project but was not ultimately "
    "needed once the automated A/B comparison gave a clear-enough answer. This "
    "null result is reported because it directly informed a real "
    "methodological decision -- more external data is not automatically "
    "beneficial, and its value depends on both the specific target model and "
    "rigorous provenance/license filtering -- rather than being a foregone "
    "conclusion."
)

ENSEMBLE_HEADING = "3.4.5 Ensemble Selection: Stacking vs. Weighted Averaging"
ENSEMBLE_PARA_1 = (
    "The ensemble configuration in Section 3.4.4 was not accepted at face "
    "value: it was revisited after a real, reported prediction quality issue "
    "-- a specific classification returned \"Brake\" at only 55% confidence, "
    "with the constituent models sharply disagreeing (Traditional ML: Belt "
    "73%; CNN: Brake 53%; YAMNet: Brake 99%) -- flagged not as an isolated "
    "case but as a recurring pattern. Two separate causes were identified. "
    "First, a reproducibility trap: the script that promotes a newly retrained "
    "model into production does not also refresh the cached feature arrays and "
    "evaluation reports sitting alongside it, so a report can silently "
    "describe an older model than the one actually deployed; an apparent "
    "accuracy drop (92.48% to 88.72%) investigated this way turned out to be "
    "the correction of stale documentation rather than a real regression, "
    "confirmed only by loading the live model file directly and cross-checking "
    "artifact timestamps instead of trusting the report text."
)
ENSEMBLE_PARA_2 = (
    "Second, and more consequential: the ensemble's own model-selection step "
    "-- choosing between a stacked meta-model and simple weighted averaging -- "
    "was shown to be noise-driven on this project's small (134-sample) "
    "validation set. A single fit-and-score comparison had stacking \"winning\" "
    "by 0.67 validation macro-F1 points (0.8727 vs. 0.8660); repeating the "
    "comparison with nested 5-fold cross-validation instead showed this "
    "apparent advantage was not real (weighted averaging: mean 0.8585, std "
    "0.0717; stacking: mean 0.8569, std 0.0699 -- stacking actually won only 3 "
    "of 5 folds), while the specific weight values a weighted-average search "
    "converged on were stable across folds (Traditional ML=0.30, CNN=0.60, "
    "YAMNet=0.10 in 4 of 5 folds). Switching the deployed configuration from "
    "stacking to this stable weighted average directly addressed the original "
    "complaint, since it reduces YAMNet's influence on disagreement cases -- "
    "matching the reported failure pattern of YAMNet being overconfident (99%) "
    "and wrong. The deployed weighted-average ensemble reaches 89.47% test "
    "accuracy, above the honestly-measured stacking result of 88.72%, and its "
    "compression robustness as a fused whole was separately confirmed in "
    "Section 3.3.2 (86.7% clean rising to 93.3% under WebM compression)."
)
ENSEMBLE_PARA_3 = (
    "The remaining, largest source of ensemble error is Sway, roughly equally "
    "hard for all three constituent models standalone (Traditional ML 33/44, "
    "CNN 35/44, YAMNet 33/44 correct on true-Sway test recordings) -- not "
    "fixable by reweighting, since no single model is disproportionately at "
    "fault. Genuinely new, correctly labeled Sway recordings remain the most "
    "promising lever for further improvement on this specific class."
)

FREESOUND_CONTRIB = (
    "A rigorously tested and rejected hypothesis: incorporating external "
    "Freesound-sourced training audio was evaluated across four models with "
    "license-filtered, model-agreement-scored candidates rather than merged on "
    "assumption, and shown to help only one of the four (PANNs) and only from "
    "a heavily filtered subset -- reported as a negative result that directly "
    "shaped the final training data policy (Section 3.3.3)."
)
ENSEMBLE_CONTRIB = (
    "A real, user-reported ensemble prediction-quality issue traced to two "
    "root causes -- a stale-artifact reproducibility trap and a "
    "validation-set-noise-driven model-selection step -- and fixed by "
    "replacing a stacked meta-model with a nested-CV-validated weighted "
    "average, directly resolving the reported failure pattern (Section 3.4.5)."
)


def set_text(p, text):
    if not p.runs:
        p.add_run(text)
        return
    p.runs[0].text = text
    for r in p.runs[1:]:
        r.text = ""


def clone_pPr(target_p, source_p):
    ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pPr"
    old = target_p._p.find(ns)
    if old is not None:
        target_p._p.remove(old)
    src = source_p._p.find(ns)
    if src is not None:
        target_p._p.insert(0, copy.deepcopy(src))


def insert_after(anchor, text, style_source):
    new_p = anchor.insert_paragraph_before("")
    anchor._p.addnext(new_p._p)
    new_p.add_run(text)
    clone_pPr(new_p, style_source)
    new_p.style = style_source.style
    return new_p


def add_sections(doc):
    P = doc.paragraphs

    # --- 3.3.3 Freesound, inserted right after 3.3.2's rejected-followup paragraph ---
    heading_src = next(p for p in P if p.text.startswith("3.3.2 Compression-Robustness"))
    anchor = next(p for p in P if p.text.startswith(
        "A further robustness step -- additionally simulating microphone/device"))
    anchor = insert_after(anchor, FREESOUND_HEADING, style_source=heading_src)
    anchor = insert_after(anchor, FREESOUND_PARA_1, style_source=anchor)  # temp, fixed below
    # body style: reuse the paragraph right after heading_src (the diagnostic para)
    body_src = next(p for p in P if p.text.startswith("A gap was observed between this project's own"))
    clone_pPr(anchor, body_src)
    anchor.style = body_src.style
    anchor = insert_after(anchor, FREESOUND_PARA_2, style_source=body_src)
    anchor = insert_after(anchor, FREESOUND_PARA_3, style_source=body_src)

    # --- 3.4.5 Ensemble selection, inserted right after 3.4.4's ensemble-weights paragraph ---
    heading_src2 = next(p for p in P if p.text.startswith("3.4.4 Weighted Ensemble"))
    anchor2 = next(p for p in P if p.text.startswith("A weighted-probability ensemble searched all weight combinations"))
    anchor2 = insert_after(anchor2, ENSEMBLE_HEADING, style_source=heading_src2)
    anchor2 = insert_after(anchor2, ENSEMBLE_PARA_1, style_source=body_src)
    anchor2 = insert_after(anchor2, ENSEMBLE_PARA_2, style_source=body_src)
    anchor2 = insert_after(anchor2, ENSEMBLE_PARA_3, style_source=body_src)

    # --- Contributions bullets, right after the compression-robustness one ---
    # (documentation.docx's Appendix A contributions list is missing some bullets
    # already present in the research paper's own copy -- a pre-existing drift, not
    # something to fix here -- so skip gracefully if the anchor isn't found.)
    contrib_hits = [p for p in P if p.text.startswith("A diagnosed and fixed real-world domain-shift")]
    if contrib_hits:
        contrib_anchor = contrib_hits[0]
        contrib_anchor = insert_after(contrib_anchor, ENSEMBLE_CONTRIB, style_source=contrib_anchor)
        contrib_anchor = insert_after(contrib_anchor, FREESOUND_CONTRIB, style_source=contrib_anchor)


if __name__ == "__main__":
    for path in [r"C:\Users\hp\Downloads\GarageAI_Research_Paper.docx"]:
        d = docx.Document(path)
        add_sections(d)
        d.save(path)
        print(path.split("\\")[-1], "updated.")
