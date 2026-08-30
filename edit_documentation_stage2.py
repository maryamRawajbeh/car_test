# -*- coding: utf-8 -*-
"""Stage 2: remaining Appendix A text fixes (Table 8 caption, Limitations,
Conclusion -- exact same substitutions as the research paper) plus the two
independently-worded summary-chapter mentions of the old 40-run study."""
import docx

SRC = r"C:\Users\hp\Downloads\GarageAI_Documentation.docx"


def set_text(paragraph, new_text):
    if not paragraph.runs:
        paragraph.add_run(new_text)
        return
    paragraph.runs[0].text = new_text
    for r in paragraph.runs[1:]:
        r.text = ""


def find_all(paragraphs, startswith):
    return [p for p in paragraphs if p.text.startswith(startswith)]


def main():
    d = docx.Document(SRC)
    P = d.paragraphs
    fixes = 0

    # --- Ch 1.5 summary ---
    for p in find_all(P, "This chapter summarizes the machine-learning work"):
        old = p.runs[0].text
        new = old.replace("a 40-run augmentation/seed-stability study",
                           "a 90-run augmentation/seed-stability study covering all nine models")
        assert new != old
        set_text(p, new)
        fixes += 1

    # --- Ch 6.1 summary ---
    for p in find_all(P, "Two complementary levels of testing were used"):
        old = p.runs[0].text
        new = old.replace(
            "plus a dedicated 40-run seed-stability study for four of the transfer-learning backbones",
            "plus a dedicated 90-run seed-stability study covering all nine individual models"
        )
        assert new != old
        set_text(p, new)
        fixes += 1

    # --- Appendix A: Table 8 caption ---
    for p in find_all(P, "Table 8. Full results of the 40-run stability study"):
        set_text(p, "Table 8. Full results of the 90-run comprehensive stability "
                 "study: nine models, five random seeds each, with and without "
                 "waveform augmentation. Every model trained on the full "
                 "626-recording training set; validation and test partitions held "
                 "fixed throughout.")
        fixes += 1

    # --- Appendix A: Limitations "two important comparisons" ---
    for p in find_all(P, "Two important comparisons remain incomplete"):
        set_text(p, "This was resolved in a 2026-08-27 follow-up: a comprehensive "
                 "90-run stability study now covers all nine individual models (not "
                 "just four), each trained on the full 626-recording training set "
                 "(not a 201-recording subset for three of them), removing both gaps "
                 "at once (Section 6.3, Tables 7-8). A new gap surfaced by the same "
                 "follow-up check is more concerning: the currently deployed SVM and "
                 "YAMNet models show a large train-vs-test accuracy gap when "
                 "evaluated on their own literal training recordings (SVM: 99.7% "
                 "train vs 86.5% test, a 13-point gap; YAMNet: 100.0% train vs 78.2% "
                 "test, a 22-point gap; CNN's gap is smaller, about 1 point on this "
                 "metric) -- both models have enough capacity to nearly memorize "
                 "their ~626 training examples, which was not previously measured "
                 "or discussed. The reported test-set numbers remain honest (the "
                 "test set was never used for fitting), but this gap suggests "
                 "real-world accuracy on data more varied than the current "
                 "133-recording test set could plausibly be lower than reported, "
                 "particularly for YAMNet.")
        fixes += 1

    # --- Appendix A: Conclusion ---
    for p in find_all(P, "This research evaluated an audio-based automotive fault classification system across ten configurations"):
        old = p.text
        marker = ("A dedicated 40-run stability study further showed that waveform "
                   "augmentation did not benefit any of four additional "
                   "transfer-learning backbones tested, and that seed-to-seed "
                   "variation alone can shift test accuracy by several percentage "
                   "points on this study's fixed 133-recording test set — a "
                   "methodological caveat that should temper strong claims about "
                   "small differences between models.")
        replacement = ("A follow-up 90-run stability study across all nine "
            "individual models further showed that waveform augmentation did not "
            "statistically benefit any of them (three -- Traditional ML, AST, and "
            "PaSST -- were measurably hurt by it), that CLAP and PaSST also exceed "
            "the traditional baseline once trained on the same full recording set "
            "as AST and EfficientAT, and that seed-to-seed variation alone can "
            "shift test accuracy by up to 8 points for frozen embeddings and 30 "
            "points in one CNN outlier run -- a methodological caveat that should "
            "temper strong claims about small differences between models. A "
            "subsequent overfitting check also found a large gap between training "
            "and held-out accuracy for the deployed SVM (99.7% vs 86.5%) and "
            "YAMNet (100.0% vs 78.2%) models, not previously measured in this "
            "study, suggesting real-world accuracy on more varied data than the "
            "current 133-recording test set could be lower than reported for those "
            "two models specifically.")
        if marker in old:
            new = old.replace(marker, replacement)
        else:
            new = old + " " + replacement
        set_text(p, new)
        fixes += 1

    d.save(SRC)
    print(f"Stage 2 done. {fixes} paragraph(s) fixed.")


if __name__ == "__main__":
    main()
