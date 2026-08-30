# -*- coding: utf-8 -*-
"""Adds the reliability-engineering additions discovered during the 2026-08-27
git-diff review of garageai-audio-analysis/garageai-frontend (concurrency cap,
temp-file race fix, ffmpeg timeout, AST's per-request loading, strictPort) to
the "Integration Notes and Engineering Lessons" section of both documents --
research paper (once) and documentation (Ch5.5 summary + Appendix A copy)."""
import copy
import docx

NEW_BULLETS = [
    "A per-request concurrency cap (asyncio.Semaphore, PREDICT_MAX_CONCURRENCY=4) "
    "was added after realizing asyncio.wait_for's timeout only stops the caller "
    "from waiting -- the underlying thread pool worker keeps running the actual "
    "ffmpeg/inference work to completion regardless, since asyncio cannot "
    "force-cancel work already handed to a thread. Without a cap, enough "
    "overlapping slow requests could exhaust the shared thread pool and starve "
    "unrelated work in the same process.",
    "Temp-file cleanup was moved from an outer `finally` block into the worker "
    "thread that actually created the files, after discovering a race: when a "
    "request timed out, the outer handler would try to delete the raw/converted "
    "audio files while the still-running background thread was mid-use of them "
    "-- on Windows, deleting a file that is still open can raise PermissionError, "
    "silently turning a clean 504 timeout into an unhandled 500.",
    "A 30-second subprocess timeout was added to the ffmpeg conversion step "
    "itself, independent of the overall request timeout, so a hung ffmpeg "
    "process cannot stall a worker thread indefinitely.",
    "AST is deliberately never kept resident in the live service (unlike PANNs "
    "and EfficientAT) -- loading it dropped free RAM as low as 62MB on the "
    "6GB-RAM development machine. It is loaded fresh per request instead, only "
    "when a user opts in via the \"compare with an extra model\" selector, with "
    "a longer request timeout (90s backend / 110s frontend) to cover the extra "
    "load time; CLAP, PaSST, and BEATs follow the same on-demand pattern.",
    "The frontend's dev server was set to fail loudly on a port conflict "
    "(Vite's strictPort: true) rather than silently moving to the next free "
    "port -- the backend's CORS configuration only allows the expected port, so "
    "a silent port shift previously surfaced as a confusing \"could not reach "
    "the server\" error instead of an obvious \"port already in use\" one.",
]


def add_bullets_after(paragraphs_list_doc, anchor_text_startswith, doc):
    hits = [p for p in doc.paragraphs if p.text.startswith(anchor_text_startswith)]
    assert len(hits) == 1, f"expected 1 match for {anchor_text_startswith!r}, got {len(hits)}"
    anchor = hits[0]
    template = anchor._p
    for text in NEW_BULLETS:
        new_p_elm = copy.deepcopy(template)
        template.addnext(new_p_elm)
        template = new_p_elm
        new_p = docx.text.paragraph.Paragraph(new_p_elm, anchor._parent)
        for r in list(new_p.runs):
            r.text = ""
        if new_p.runs:
            new_p.runs[0].text = text
        else:
            new_p.add_run(text)


# --- Research paper: one "Integration Notes" section ---
d1 = docx.Document(r"C:\Users\hp\Downloads\GarageAI_Research_Paper.docx")
add_bullets_after(d1.paragraphs, "All code comments, error messages, Swagger descriptions, and interface text were kept in English.", d1)
d1.save(r"C:\Users\hp\Downloads\GarageAI_Research_Paper.docx")
print("Research paper updated.")

# --- Documentation: two copies (Ch5.5 summary + Appendix A) ---
d2 = docx.Document(r"C:\Users\hp\Downloads\GarageAI_Documentation.docx")
targets = [p for p in d2.paragraphs if p.text.startswith("All code comments, error messages, Swagger descriptions, and interface text")]
assert len(targets) == 2, f"expected 2 matches in documentation.docx, got {len(targets)}"
for anchor in targets:
    template = anchor._p
    for text in NEW_BULLETS:
        new_p_elm = copy.deepcopy(template)
        template.addnext(new_p_elm)
        template = new_p_elm
        new_p = docx.text.paragraph.Paragraph(new_p_elm, anchor._parent)
        for r in list(new_p.runs):
            r.text = ""
        if new_p.runs:
            new_p.runs[0].text = text
        else:
            new_p.add_run(text)
d2.save(r"C:\Users\hp\Downloads\GarageAI_Documentation.docx")
print("Documentation updated (both sections).")
