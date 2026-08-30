# -*- coding: utf-8 -*-
r"""
Append "Appendix C: User Interface Screenshots" to GarageAI_Documentation.docx
-- a dark-theme gallery of the running frontend, captured 2026-08-31 against
the full local stack (frontend 5173 + Node gateway 5000 + FastAPI inference
service 8001) via Playwright. Grouped: public/auth pages, the Home analysis
flow (incl. record / result / severity / low-confidence / "other" states),
Ask GarageAI chat, and History / Settings.
"""
import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

DOC = r"C:\Users\hp\Downloads\GarageAI_Documentation.docx"
SHOT_DIR = r"C:\Users\hp\AppData\Local\Temp\claude\c--Users-hp-Desktop-car-test\f64fe89b-64da-445d-b0e5-e0e0c3213a35\scratchpad\ui"

INTRO = (
    "This appendix shows the deployed GarageAI web interface (React / Vite frontend, "
    "port 5173) running against the live Node gateway and FastAPI inference service. "
    "All captures use the application's Dark theme at a desktop viewport and were taken "
    "on 2026-08-31 with real model inference, except three states (\u201cNo Clear Match\u201d, "
    "the \u201cother\u201d class, and the seeded chat answer) which use scripted API responses to "
    "reproduce a specific screen deterministically."
)

# (subsection title, [(file, width_in, caption), ...])
GROUPS = [
    ("C.1 Public and Authentication Pages", [
        ("ui_01_login.png", 6.0, "Figure C.1. The login page."),
        ("ui_02_signup.png", 6.0, "Figure C.2. The sign-up page."),
        ("ui_03_forgot_password.png", 6.0, "Figure C.3. The forgot-password page."),
        ("ui_04_about.png", 5.2, "Figure C.4. The About page."),
    ]),
    ("C.2 Home \u2014 Sound Analysis Flow", [
        ("ui_05_home_empty.png", 6.0, "Figure C.5. The Home page before any audio is selected."),
        ("ui_06_home_extra_model.png", 6.0, "Figure C.6. Home with an optional extra comparison model (AST) selected; a note warns of the added latency."),
        ("ui_07_home_recording.png", 6.0, "Figure C.7. Recording from the browser microphone: live level meter and the Stop Recording control."),
        ("ui_08_home_result.png", 6.0, "Figure C.8. A completed analysis \u2014 predicted class, overall confidence, per-class confidence bars, and each model's individual vote."),
        ("ui_09_severity_dialog.png", 6.0, "Figure C.9. The severity checklist dialog (opened from Assess Severity) \u2014 quick yes/no questions asked in Arabic."),
        ("ui_10_home_no_match.png", 6.0, "Figure C.10. A low-confidence result rendered as \u201cNo Clear Match\u201d, with the closest guess shown."),
        ("ui_11_home_other.png", 6.0, "Figure C.11. A result classified as \u201cother\u201d \u2014 the audio matches none of belt, brake, or sway; the per-model breakdown is omitted."),
    ]),
    ("C.3 Ask GarageAI (Chat)", [
        ("ui_12_chat_empty.png", 6.0, "Figure C.12. The Ask GarageAI chat before any message is sent."),
        ("ui_13_chat_conversation.png", 6.0, "Figure C.13. A chat exchange: a free-text question and the knowledge-base-grounded answer."),
    ]),
    ("C.4 History and Settings", [
        ("ui_14_history.png", 6.0, "Figure C.14. The History page listing past analyses with full result detail; searchable and filterable by issue."),
        ("ui_15_settings.png", 6.0, "Figure C.15. The Settings page \u2014 personal information, account security, appearance (theme), and notification preferences."),
    ]),
]


def style_like(d, *prefixes):
    """Return the .style of the first existing paragraph whose text starts with
    one of the given prefixes -- more reliable than d.styles[name] for docs whose
    heading styles are not registered under the canonical English names."""
    for p in d.paragraphs:
        t = p.text.strip()
        for pre in prefixes:
            if t.startswith(pre):
                return p.style
    return None


def main():
    d = docx.Document(DOC)
    h1 = style_like(d, "Appendix A", "Appendix B", "9. References")
    h3 = style_like(d, "3.4.5 Ensemble Selection", "B.1 ", "7.1 Recording")
    cap_style = None
    for p in d.paragraphs:
        if p.text.strip().startswith("Figure 7."):
            cap_style = p.style
            break

    d.add_page_break()
    h = d.add_paragraph("Appendix C: User Interface Screenshots")
    if h1 is not None:
        h.style = h1
    d.add_paragraph(INTRO)

    for title, items in GROUPS:
        hs = d.add_paragraph(title)
        if h3 is not None:
            hs.style = h3
        for fname, width_in, caption in items:
            pic_p = d.add_paragraph()
            pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pic_p.add_run().add_picture(f"{SHOT_DIR}\\{fname}", width=Inches(width_in))
            cap = d.add_paragraph(caption)
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if cap_style is not None:
                cap.style = cap_style
            else:
                cap.runs[0].italic = True
                cap.runs[0].font.size = Pt(9)

    d.save(DOC)
    print("saved", DOC, "-- inline shapes now:", len(d.inline_shapes))


if __name__ == "__main__":
    main()
