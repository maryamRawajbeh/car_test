# -*- coding: utf-8 -*-
r"""
Two corrections to the ensemble-retune edits:

1. Section 3.4.6 (heading + body) was inserted right after the Section 3.4.5
   sway paragraph, which put it BEFORE "Figure 3. Optimal ensemble weights
   selected via exhaustive validation-set search." and its image -- a figure
   that belongs to the ORIGINAL weight search (3.4.4/3.4.5), not the rejected
   retune. Move the 3.4.6 heading+body to immediately before "3.5 Training
   Process" so Figure 3 stays under 3.4.5 and 3.4.6 reads as its own thing.

2. Discussion 7.2 still said the deployed weights "was not re-run against"
   the extended comparison -- now stale, since 3.4.6 is exactly that re-run.
   Soften to point at 3.4.6.
"""
import docx

FILES = [r"C:\Users\hp\Downloads\GarageAI_Research_Paper.docx",
         r"C:\Users\hp\Downloads\GarageAI_Documentation.docx"]

STALE = ("a configuration that predates the extended comparison and was not "
         "re-run against it")
FIXED = ("a configuration that predates the extended comparison; Section 3.4.6 "
         "describes re-running the weight search against the full model set")


def set_text(p, txt):
    if not p.runs:
        p.add_run(txt); return
    p.runs[0].text = txt
    for r in p.runs[1:]:
        r.text = ""


def main():
    for fn in FILES:
        d = docx.Document(fn)
        P = d.paragraphs

        h346 = next(p for p in P if p.text.strip().startswith("3.4.6 A Post-Hoc"))
        body346 = h346._p.getnext()  # the 3.4.6 body paragraph
        assert body346.text if hasattr(body346, "text") else True
        h35 = next(p for p in P if p.text.strip().startswith("3.5 Training Process"))

        # detach the two 3.4.6 paragraphs and re-insert them right before 3.5
        h346_el, body346_el = h346._p, body346
        h346_el.getparent().remove(h346_el)
        body346_el.getparent().remove(body346_el)
        h35._p.addprevious(h346_el)
        h346_el.addnext(body346_el)

        # fix the stale 7.2 clause
        disc = next((p for p in d.paragraphs if STALE in p.text), None)
        assert disc is not None, f"stale 7.2 clause not found in {fn}"
        set_text(disc, disc.text.replace(STALE, FIXED))

        d.save(fn)
        print("fixed", fn)


if __name__ == "__main__":
    main()
