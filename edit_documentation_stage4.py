# -*- coding: utf-8 -*-
"""Stage 4: replace Appendix A's old 4-model/40-run summary and raw tables with
the comprehensive 9-model/90-run data (same as done for the research paper)."""
import copy
import docx
import pandas as pd
import numpy as np

SRC = r"C:\Users\hp\Downloads\GarageAI_Documentation.docx"


def resize_table_rows(table, n_data_rows_needed):
    current_data_rows = len(table.rows) - 1
    if current_data_rows < n_data_rows_needed:
        template_row = table.rows[-1]._tr
        for _ in range(n_data_rows_needed - current_data_rows):
            new_row = copy.deepcopy(template_row)
            template_row.addnext(new_row)
            template_row = new_row
    elif current_data_rows > n_data_rows_needed:
        for row in table.rows[n_data_rows_needed + 1:][::-1]:
            row._tr.getparent().remove(row._tr)


def build_summary_rows():
    name_map = {"yamnet": "YAMNet", "panns": "PANNs/CNN14", "ast": "AST",
                "clap": "CLAP", "passt": "PaSST", "beats": "BEATs"}
    frozen = pd.read_csv("processed_data/augmentation_variance_study_full_data.csv")
    tml = pd.read_csv("processed_data/traditional_ml_augmentation_variance.csv")
    cnn = pd.read_csv("processed_data/cnn_augmentation_variance.csv")
    eat = pd.read_csv("processed_data/efficientat_augmentation_study.csv")

    stats = {}
    for (model, aug), g in frozen.groupby(["model", "augment"]):
        stats.setdefault(name_map[model], {})[bool(aug)] = g["test_accuracy"].values * 100
    for cond, g in tml.groupby("condition"):
        stats.setdefault("Traditional ML (SVM/XGBoost)", {})[cond == "with_augmentation"] = \
            g["test_accuracy"].values * 100
    for cond, g in cnn.groupby("condition"):
        stats.setdefault("CNN", {})[cond == "with_augmentation"] = g["test_accuracy"].values * 100
    for aug, g in eat.groupby("augment"):
        stats.setdefault("EfficientAT (fine-tuned)", {})[bool(aug)] = g["test_accuracy"].values * 100

    order = ["YAMNet", "PANNs/CNN14", "CNN", "BEATs", "PaSST", "CLAP",
             "Traditional ML (SVM/XGBoost)", "AST", "EfficientAT (fine-tuned)"]
    rows = []
    for name in order:
        no_aug = stats[name][False]
        with_aug = stats[name][True]
        m0, s0 = no_aug.mean(), no_aug.std()
        m1, s1 = with_aug.mean(), with_aug.std()
        gap = m1 - m0
        pooled = np.sqrt((s0 ** 2 + s1 ** 2) / 2)
        if pooled > 0 and abs(gap) > 1.5 * pooled:
            verdict = "Hurts" if gap < 0 else "Helps"
        else:
            verdict = "No effect (noise)"
        rows.append([name, f"{m0:.1f}% (+/-{s0:.1f})", f"{m1:.1f}% (+/-{s1:.1f})",
                     f"{gap:+.1f} pts -- {verdict}"])
    return rows


def build_raw_rows():
    name_map = {"yamnet": "YAMNet", "panns": "PANNs/CNN14", "ast": "AST",
                "clap": "CLAP", "passt": "PaSST", "beats": "BEATs"}
    frozen = pd.read_csv("processed_data/augmentation_variance_study_full_data.csv")
    tml = pd.read_csv("processed_data/traditional_ml_augmentation_variance.csv")
    cnn = pd.read_csv("processed_data/cnn_augmentation_variance.csv")
    eat = pd.read_csv("processed_data/efficientat_augmentation_study.csv")

    order = ["YAMNet", "PANNs/CNN14", "CNN", "BEATs", "PaSST", "CLAP",
              "Traditional ML (SVM/XGBoost)", "AST", "EfficientAT (fine-tuned)"]
    by_model = {name: [] for name in order}

    for _, r in frozen.iterrows():
        name = name_map[r["model"]]
        n = 2504 if r["augment"] else 626
        by_model[name].append([name, str(int(r["seed"])), "Yes" if r["augment"] else "No",
                                str(n), f"{r['test_accuracy']*100:.2f}%", f"{r['test_macro_f1']:.4f}"])
    for _, r in tml.iterrows():
        aug = r["condition"] == "with_augmentation"
        n = 1872 if aug else 624
        by_model["Traditional ML (SVM/XGBoost)"].append([
            "Traditional ML (SVM/XGBoost)", str(int(r["seed"])), "Yes" if aug else "No",
            str(n), f"{r['test_accuracy']*100:.2f}%", f"{r['test_macro_f1']:.4f}"])
    for _, r in cnn.iterrows():
        aug = r["condition"] == "with_augmentation"
        n = 1872 if aug else 624
        by_model["CNN"].append(["CNN", str(int(r["seed"])), "Yes" if aug else "No",
                                 str(n), f"{r['test_accuracy']*100:.2f}%", f"{r['test_macro_f1']:.4f}"])
    for _, r in eat.iterrows():
        by_model["EfficientAT (fine-tuned)"].append([
            "EfficientAT (fine-tuned)", str(int(r["seed"])), "Yes" if r["augment"] else "No",
            "626 (online aug.)" if r["augment"] else "626",
            f"{r['test_accuracy']*100:.2f}%", f"{r['test_macro_f1']:.4f}"])

    rows = []
    for name in order:
        rows.extend(sorted(by_model[name], key=lambda x: (int(x[1]), x[2])))
    return rows


def main():
    d = docx.Document(SRC)
    T = d.tables

    t_summary = T[12]
    assert len(t_summary.rows) == 5 and len(t_summary.columns) == 4
    summary_rows = build_summary_rows()
    resize_table_rows(t_summary, len(summary_rows))
    header = ["Model", "Mean Accuracy (no augmentation)", "Mean Accuracy (with augmentation)", "Effect"]
    for c, text in enumerate(header):
        t_summary.cell(0, c).text = text
    for r, row in enumerate(summary_rows):
        for c, val in enumerate(row):
            t_summary.cell(r + 1, c).text = val
    print(f"Summary table resized to {len(summary_rows)} data rows.")

    t_raw = T[13]
    assert len(t_raw.rows) == 41 and len(t_raw.columns) == 6
    raw_rows = build_raw_rows()
    resize_table_rows(t_raw, len(raw_rows))
    header = ["Model", "Seed", "Augmentation", "Train Size", "Test Accuracy", "Test Macro-F1"]
    for c, text in enumerate(header):
        t_raw.cell(0, c).text = text
    for r, row in enumerate(raw_rows):
        for c, val in enumerate(row):
            t_raw.cell(r + 1, c).text = val
    print(f"Raw table resized to {len(raw_rows)} data rows.")

    d.save(SRC)
    print("Stage 4 done.")


if __name__ == "__main__":
    main()
